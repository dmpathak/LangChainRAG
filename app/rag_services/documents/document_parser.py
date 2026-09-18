"""Convert uploaded files into LangChain documents."""

import hashlib
import os

import fitz
import pandas as pd
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import CHUNK_OVERLAP, CHUNK_SIZE, OCR_ENABLED

try:
    import pytesseract
    from PIL import Image
except ImportError:
    pytesseract = None
    Image = None


TEXT_SPLITTER = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
)


class DocumentProcessor:
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".csv", ".xls", ".xlsx"}

    def process(self, file, metadata=None):
        metadata = metadata or {}
        extension = os.path.splitext(file.filename)[1].lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {extension}")
        if extension == ".pdf":
            return self._process_pdf(file.file, metadata)
        if extension == ".docx":
            return self._process_docx(file.file, metadata)
        if extension == ".csv":
            return self._process_csv(file.file, metadata)
        return self._process_excel(file.file, metadata)

    def _process_pdf(self, file, metadata):
        file.seek(0)
        pdf = fitz.open(stream=file.read(), filetype="pdf")
        documents = []
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text()
            if not text.strip() and OCR_ENABLED:
                text = self._extract_text_with_ocr(page)
            if not text.strip():
                continue
            page_metadata = {**metadata, "page_number": page_number}
            documents.extend(self._create_chunked_documents(text, page_metadata))
        if not documents:
            raise ValueError(
                "Unable to extract text from PDF. For scanned PDFs, install "
                "pytesseract, Pillow, and the Tesseract system package."
            )
        return documents

    @staticmethod
    def _extract_text_with_ocr(page):
        if pytesseract is None or Image is None:
            return ""
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
        return pytesseract.image_to_string(image)

    def _process_docx(self, file, metadata):
        file.seek(0)
        document = DocxDocument(file)
        text = "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        )
        if not text:
            raise ValueError("Unable to extract text from DOCX.")
        return self._create_chunked_documents(text, metadata)

    def _process_csv(self, file, metadata):
        file.seek(0)
        try:
            dataframe = pd.read_csv(file)
        except UnicodeDecodeError:
            file.seek(0)
            dataframe = pd.read_csv(file, encoding="latin-1")
        return self._create_row_documents(dataframe, metadata)

    def _process_excel(self, file, metadata):
        file.seek(0)
        documents = []
        for sheet_name, dataframe in pd.read_excel(file, sheet_name=None).items():
            documents.extend(
                self._create_row_documents(
                    dataframe,
                    {**metadata, "sheet_name": sheet_name},
                )
            )
        return documents

    def _create_chunked_documents(self, text, metadata):
        documents = TEXT_SPLITTER.split_documents(
            [Document(page_content=text, metadata=metadata)]
        )
        for index, document in enumerate(documents):
            document.metadata["chunk_index"] = index
            document.metadata["document_id"] = self._document_id(document)
        return documents

    def _create_row_documents(self, dataframe, metadata):
        documents = []
        for row_number, (_, row) in enumerate(dataframe.iterrows(), start=1):
            if row.isna().all():
                continue
            values = {
                self._normalize_column_name(column): self._normalize_metadata_value(value)
                for column, value in row.items()
            }
            document_metadata = {**metadata, **values, "row_number": row_number}
            content = "\n".join(f"{key}: {value}" for key, value in values.items())
            document = Document(page_content=content, metadata=document_metadata)
            document.metadata["document_id"] = self._document_id(document)
            documents.append(document)
        return documents

    @staticmethod
    def _normalize_metadata_value(value):
        if pd.isna(value):
            return None
        if hasattr(value, "item"):
            value = value.item()
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value

    @staticmethod
    def _normalize_column_name(value):
        return "_".join(str(value).strip().lower().replace("-", " ").split())

    @staticmethod
    def _document_id(document):
        identity = "|".join(
            str(document.metadata.get(key, ""))
            for key in ("file_name", "sheet_name", "page_number", "row_number", "chunk_index")
        ) + "|" + document.page_content
        return hashlib.sha256(identity.encode("utf-8")).hexdigest()
