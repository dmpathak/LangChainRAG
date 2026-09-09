"""Convert uploaded files into LangChain documents."""

import hashlib
import os

import fitz
import pandas as pd
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from config import CHUNK_OVERLAP, CHUNK_SIZE


TEXT_SPLITTER = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
)


class DocumentProcessor:
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".csv", ".xls", ".xlsx"}

    def process(self, file, metadata=None):
        metadata = metadata or {}
        extension = os.path.splitext(file.filename)[1].lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {extension}")

        if extension == ".pdf":
            return self._process_pdf(file.file, metadata)
        if extension == ".docx":
            return self._process_docx(file.file, metadata)
        if extension == ".csv":
            return self._process_csv(file.file, metadata)
        return self._process_excel(file.file, metadata)

    def _process_pdf(self, file, metadata):
        file.seek(0)
        pdf = fitz.open(stream=file.read(), filetype="pdf")
        documents = []

        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text()
            if not text.strip():
                continue

            page_metadata = metadata.copy()
            page_metadata["page_number"] = page_number
            page_documents = self._create_chunked_documents(text, page_metadata)
            documents.extend(page_documents)

        if not documents:
            raise ValueError("Unable to extract text from PDF.")
        return documents

    def _process_docx(self, file, metadata):
        file.seek(0)
        document = DocxDocument(file)
        paragraphs = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        text = "\n".join(paragraphs)
        if not text:
            raise ValueError("Unable to extract text from DOCX.")
        return self._create_chunked_documents(text, metadata)

    def _process_csv(self, file, metadata):
        file.seek(0)
        try:
            dataframe = pd.read_csv(file)
        except UnicodeDecodeError:
            file.seek(0)
            dataframe = pd.read_csv(file, encoding="latin-1")
        return self._create_row_documents(dataframe, metadata)

    def _process_excel(self, file, metadata):
        file.seek(0)
        sheets = pd.read_excel(file, sheet_name=None)
        documents = []

        for sheet_name, dataframe in sheets.items():
            sheet_metadata = metadata.copy()
            sheet_metadata["sheet_name"] = sheet_name
            row_documents = self._create_row_documents(dataframe, sheet_metadata)
            documents.extend(row_documents)

        return documents

    def _create_chunked_documents(self, text, metadata):
        source_document = Document(page_content=text, metadata=metadata)
        documents = TEXT_SPLITTER.split_documents([source_document])

        for index, document in enumerate(documents):
            document.metadata["chunk_index"] = index
            document.metadata["document_id"] = self._document_id(document)

        return documents

    def _create_row_documents(self, dataframe, metadata):
        documents = []

        for row_number, row_data in enumerate(dataframe.iterrows(), start=1):
            row = row_data[1]
            if row.isna().all():
                continue

            row_metadata = {}
            for column, value in row.items():
                column_name = self._normalize_column_name(column)
                row_metadata[column_name] = self._normalize_metadata_value(value)

            row_lines = []
            for column, value in row_metadata.items():
                row_lines.append(f"{column}: {value}")

            document_metadata = metadata.copy()
            document_metadata.update(row_metadata)
            document_metadata["row_number"] = row_number

            document = Document(
                page_content="\n".join(row_lines),
                metadata=document_metadata,
            )
            document.metadata["document_id"] = self._document_id(document)
            documents.append(document)

        return documents

    @staticmethod
    def _normalize_metadata_value(value):
        if pd.isna(value):
            return None
        if hasattr(value, "item"):
            value = value.item()
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value

    @staticmethod
    def _normalize_column_name(value):
        name = str(value).strip().lower().replace("-", " ")
        return "_".join(name.split())

    @staticmethod
    def _document_id(document):
        identity_parts = [
            str(document.metadata.get("file_name", "")),
            str(document.metadata.get("sheet_name", "")),
            str(document.metadata.get("page_number", "")),
            str(document.metadata.get("row_number", "")),
            str(document.metadata.get("chunk_index", "")),
            document.page_content,
        ]
        identity = "|".join(identity_parts)
        return hashlib.sha256(identity.encode("utf-8")).hexdigest()
